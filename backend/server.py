import sys
import os
from pathlib import Path

# Garante que backend/ está no sys.path para que `services.*` seja encontrado
# tanto ao rodar localmente (uvicorn backend.server:app) quanto no serverless Vercel
_backend_dir = Path(__file__).parent
if str(_backend_dir) not in sys.path:
    sys.path.insert(0, str(_backend_dir))

from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException, Query
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import logging
from pydantic import BaseModel, Field, ConfigDict, BeforeValidator
from typing import List, Optional, Annotated
from datetime import datetime, timezone, date, timedelta
from bson import ObjectId

from services.matching import match_product as _match_product
from services.xml_parser import parse_nfe_xml

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL_2']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME_2']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ── PyObjectId & BaseDocument ──────────────────────────────────────
PyObjectId = Annotated[str, BeforeValidator(str)]

class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True, arbitrary_types_allowed=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")

    @classmethod
    def from_mongo(cls, doc):
        if doc is None:
            return None
        doc['_id'] = str(doc['_id'])
        return cls(**doc)

    def to_mongo(self):
        data = self.model_dump(by_alias=True, exclude_none=True)
        if '_id' in data and data['_id']:
            data['_id'] = ObjectId(data['_id'])
        else:
            data.pop('_id', None)
        return data

# ── Models ─────────────────────────────────────────────────────────
class Produto(BaseDocument):
    codigo: str
    descricao: str
    ean: Optional[str] = None
    unidade: str = "UN"
    preco: float = 0.0
    categoria: Optional[str] = None
    ativo: bool = True
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class ProdutoCreate(BaseModel):
    codigo: str
    descricao: str
    ean: Optional[str] = None
    unidade: str = "UN"
    preco: float = 0.0
    categoria: Optional[str] = None

class Fornecedor(BaseDocument):
    cnpj: str
    nome: str
    contato: Optional[str] = None
    email: Optional[str] = None
    telefone: Optional[str] = None
    ativo: bool = True
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class FornecedorCreate(BaseModel):
    cnpj: str
    nome: str
    contato: Optional[str] = None
    email: Optional[str] = None
    telefone: Optional[str] = None

class ItemNota(BaseDocument):
    nota_id: str
    numero_item: int
    cprod: str = ""
    ean: Optional[str] = None
    descricao_nfe: str
    ncm: Optional[str] = None
    cfop: Optional[str] = None
    quantidade: float
    unidade: str = "UN"
    valor_unitario: float = 0.0
    valor_total: float = 0.0
    produto_interno_id: Optional[str] = None
    produto_interno_codigo: Optional[str] = None
    produto_interno_descricao: Optional[str] = None
    metodo_identificacao: Optional[str] = None
    confianca: float = 0.0
    sugestoes: Optional[list] = None
    criterios: Optional[list] = None
    status: str = "pendente"
    quantidade_conferida: float = 0.0
    justificativa_divergencia: Optional[str] = None
    ignorado: bool = False

class Nota(BaseDocument):
    chave: Optional[str] = None
    numero: Optional[str] = None
    serie: Optional[str] = None
    data_emissao: Optional[str] = None
    fornecedor_cnpj: Optional[str] = None
    fornecedor_nome: Optional[str] = None
    valor_total: float = 0.0
    status: str = "pendente"
    total_itens: int = 0
    itens_identificados: int = 0
    itens_conferidos: int = 0
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    conferencia_inicio: Optional[str] = None
    conferencia_fim: Optional[str] = None
    operador_conferencia: Optional[str] = None
    relatorio_salvo: Optional[dict] = None

class EquivalenciaProduto(BaseDocument):
    fornecedor_cnpj: str
    fornecedor_nome: str
    codigo_fornecedor: str
    ean: Optional[str] = None
    descricao_nfe: str
    produto_interno_id: str
    produto_interno_codigo: str
    produto_interno_descricao: str
    usuario: str = "operador"
    utilizacoes: int = 1
    ultima_utilizacao: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    confianca: float = 100.0
    origem_vinculo: str = "manual"  # ean, codigo_fornecedor, similaridade, manual
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class HistoricoAprendizado(BaseDocument):
    fornecedor_cnpj: str
    fornecedor_nome: str = ""
    codigo_fornecedor: str
    ean: Optional[str] = None
    descricao_nfe: str
    produto_interno_id: str
    produto_interno_codigo: str
    produto_interno_descricao: str
    usuario: str = "operador"
    confianca: float = 100.0
    metodo_vinculo: str = "manual"
    nota_id: Optional[str] = None
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class HistoricoLeitura(BaseDocument):
    nota_id: str
    codigo_barras: str
    produto_interno_id: Optional[str] = None
    item_nota_id: Optional[str] = None
    resultado: str
    timestamp: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())



# ── Matching wrapper (injects db + Produto model) ─────────────────
async def match_product(item_desc, item_ean, item_cprod, fornecedor_cnpj, quantidade=0, _cache=None):
    return await _match_product(db, Produto, item_desc, item_ean, item_cprod, fornecedor_cnpj, quantidade, _cache)

# ── Learning Service ───────────────────────────────────────────────
async def save_learning(item, produto, nota, metodo='manual', confianca=100.0):
    """Save learning data whenever a binding is confirmed."""
    if not nota or not nota.fornecedor_cnpj:
        return

    # Save to historico_aprendizado
    hist = HistoricoAprendizado(
        fornecedor_cnpj=nota.fornecedor_cnpj,
        fornecedor_nome=nota.fornecedor_nome or '',
        codigo_fornecedor=item.cprod,
        ean=item.ean,
        descricao_nfe=item.descricao_nfe,
        produto_interno_id=produto.id,
        produto_interno_codigo=produto.codigo,
        produto_interno_descricao=produto.descricao,
        confianca=confianca,
        metodo_vinculo=metodo,
        nota_id=item.nota_id,
    )
    await db.historico_aprendizado.insert_one(hist.to_mongo())

    # Upsert equivalencia
    existing_eq = await db.equivalencia_produtos.find_one({
        'fornecedor_cnpj': nota.fornecedor_cnpj, 'codigo_fornecedor': item.cprod
    })
    if existing_eq:
        new_confianca = min(existing_eq.get('confianca', 0) + 2, 100)
        await db.equivalencia_produtos.update_one(
            {'_id': existing_eq['_id']},
            {'$inc': {'utilizacoes': 1}, '$set': {
                'ultima_utilizacao': datetime.now(timezone.utc).isoformat(),
                'produto_interno_id': produto.id,
                'produto_interno_codigo': produto.codigo,
                'produto_interno_descricao': produto.descricao,
                'confianca': new_confianca,
            }}
        )
    else:
        eq = EquivalenciaProduto(
            fornecedor_cnpj=nota.fornecedor_cnpj,
            fornecedor_nome=nota.fornecedor_nome or '',
            codigo_fornecedor=item.cprod,
            ean=item.ean,
            descricao_nfe=item.descricao_nfe,
            produto_interno_id=produto.id,
            produto_interno_codigo=produto.codigo,
            produto_interno_descricao=produto.descricao,
            confianca=confianca,
            origem_vinculo=metodo,
        )
        await db.equivalencia_produtos.insert_one(eq.to_mongo())

    # Update product EAN if missing
    if item.ean and not produto.ean:
        await db.produtos.update_one({'_id': ObjectId(produto.id)}, {'$set': {'ean': item.ean}})

# ── API: Health ────────────────────────────────────────────────────
@api_router.get("/")
async def root():
    return {"message": "NF-e Conference System API"}

# ── API: Produtos ──────────────────────────────────────────────────
@api_router.get("/produtos")
async def list_produtos(search: Optional[str] = None):
    query = {'ativo': True}
    if search:
        query['$or'] = [
            {'codigo': {'$regex': search, '$options': 'i'}},
            {'descricao': {'$regex': search, '$options': 'i'}},
            {'ean': {'$regex': search, '$options': 'i'}},
        ]
    docs = await db.produtos.find(query).sort('codigo', 1).to_list(2000)
    return [Produto.from_mongo(d).model_dump() for d in docs]

@api_router.get("/produtos/exportar-excel")
async def exportar_produtos_excel():
    """Export all active products to an Excel spreadsheet (cod interno, descricao, EAN)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    import io

    docs = await db.produtos.find({'ativo': True}).sort('codigo', 1).to_list(10000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"

    headers = ["Codigo Interno", "Descricao", "EAN"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for row, d in enumerate(docs, 2):
        ws.cell(row=row, column=1, value=d.get('codigo', ''))
        ws.cell(row=row, column=2, value=d.get('descricao', ''))
        ws.cell(row=row, column=3, value=d.get('ean') or '')

    # Auto column widths
    for col in range(1, 4):
        max_len = len(headers[col - 1])
        for row in range(2, len(docs) + 2):
            v = ws.cell(row=row, column=col).value
            if v:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"produtos_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@api_router.post("/produtos")
async def create_produto(data: ProdutoCreate):
    existing = await db.produtos.find_one({'codigo': data.codigo, 'ativo': True})
    if existing:
        raise HTTPException(400, f"Produto com codigo {data.codigo} ja existe")
    produto = Produto(**data.model_dump())
    result = await db.produtos.insert_one(produto.to_mongo())
    produto.id = str(result.inserted_id)
    return produto.model_dump()

@api_router.put("/produtos/{produto_id}")
async def update_produto(produto_id: str, data: ProdutoCreate):
    result = await db.produtos.update_one({'_id': ObjectId(produto_id)}, {'$set': data.model_dump()})
    if result.matched_count == 0:
        raise HTTPException(404, "Produto nao encontrado")
    doc = await db.produtos.find_one({'_id': ObjectId(produto_id)})
    return Produto.from_mongo(doc).model_dump()

@api_router.delete("/produtos/{produto_id}")
async def delete_produto(produto_id: str):
    await db.produtos.update_one({'_id': ObjectId(produto_id)}, {'$set': {'ativo': False}})
    return {"ok": True}

# ── API: Fornecedores ──────────────────────────────────────────────
@api_router.get("/fornecedores")
async def list_fornecedores(search: Optional[str] = None):
    query = {'ativo': True}
    if search:
        query['$or'] = [
            {'cnpj': {'$regex': search, '$options': 'i'}},
            {'nome': {'$regex': search, '$options': 'i'}},
        ]
    docs = await db.fornecedores.find(query).sort('nome', 1).to_list(1000)
    return [Fornecedor.from_mongo(d).model_dump() for d in docs]

@api_router.post("/fornecedores")
async def create_fornecedor(data: FornecedorCreate):
    fornecedor = Fornecedor(**data.model_dump())
    result = await db.fornecedores.insert_one(fornecedor.to_mongo())
    fornecedor.id = str(result.inserted_id)
    return fornecedor.model_dump()

@api_router.put("/fornecedores/{fornecedor_id}")
async def update_fornecedor(fornecedor_id: str, data: FornecedorCreate):
    result = await db.fornecedores.update_one({'_id': ObjectId(fornecedor_id)}, {'$set': data.model_dump()})
    if result.matched_count == 0:
        raise HTTPException(404, "Fornecedor nao encontrado")
    doc = await db.fornecedores.find_one({'_id': ObjectId(fornecedor_id)})
    return Fornecedor.from_mongo(doc).model_dump()

@api_router.delete("/fornecedores/{fornecedor_id}")
async def delete_fornecedor(fornecedor_id: str):
    await db.fornecedores.update_one({'_id': ObjectId(fornecedor_id)}, {'$set': {'ativo': False}})
    return {"ok": True}

# ── API: Notas ─────────────────────────────────────────────────────
@api_router.get("/notas")
async def list_notas(status: Optional[str] = None):
    query = {}
    if status:
        query['status'] = status
    docs = await db.notas.find(query).sort('created_at', -1).to_list(1000)
    return [Nota.from_mongo(d).model_dump() for d in docs]

@api_router.get("/notas/{nota_id}")
async def get_nota(nota_id: str):
    doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    if not doc:
        raise HTTPException(404, "Nota nao encontrada")
    nota = Nota.from_mongo(doc)
    items = await db.itens_nota.find({'nota_id': nota_id}).sort('numero_item', 1).to_list(1000)
    return {'nota': nota.model_dump(), 'itens': [ItemNota.from_mongo(i).model_dump() for i in items]}

@api_router.post("/notas/importar-xml")
async def importar_xml(file: UploadFile = File(...)):
    content = await file.read()
    try:
        header, items = parse_nfe_xml(content)
    except Exception as e:
        logger.error(f"XML parse error: {e}")
        raise HTTPException(400, f"Erro ao processar XML: {str(e)}")

    if header.get('chave'):
        existing = await db.notas.find_one({'chave': header['chave']})
        if existing:
            raise HTTPException(400, f"Nota com chave {header['chave']} ja importada")

    if header.get('fornecedor_cnpj'):
        existing_forn = await db.fornecedores.find_one({'cnpj': header['fornecedor_cnpj']})
        if not existing_forn:
            forn = Fornecedor(cnpj=header['fornecedor_cnpj'], nome=header.get('fornecedor_nome', 'Desconhecido'))
            await db.fornecedores.insert_one(forn.to_mongo())

    nota = Nota(
        chave=header.get('chave'), numero=header.get('numero'), serie=header.get('serie'),
        data_emissao=header.get('data_emissao'), fornecedor_cnpj=header.get('fornecedor_cnpj'),
        fornecedor_nome=header.get('fornecedor_nome'), valor_total=header.get('valor_total', 0),
        total_itens=len(items),
    )
    nota_result = await db.notas.insert_one(nota.to_mongo())
    nota_id = str(nota_result.inserted_id)

    itens_identificados = 0
    match_cache = {}  # shared across items of this NF-e (products list, bindings, count)
    for item_data in items:
        match_result = await match_product(
            item_data['descricao_nfe'], item_data.get('ean'),
            item_data.get('cprod'), header.get('fornecedor_cnpj'),
            item_data.get('quantidade', 0),
            _cache=match_cache,
        )
        item = ItemNota(
            nota_id=nota_id, **item_data,
            produto_interno_id=match_result['produto']['id'] if match_result['produto'] else None,
            produto_interno_codigo=match_result['produto']['codigo'] if match_result['produto'] else None,
            produto_interno_descricao=match_result['produto']['descricao'] if match_result['produto'] else None,
            metodo_identificacao=match_result['metodo'],
            confianca=match_result['confianca'],
            sugestoes=match_result.get('sugestoes', []),
            criterios=match_result.get('criterios', []),
            status='identificado' if match_result['produto'] else 'pendente'
        )
        await db.itens_nota.insert_one(item.to_mongo())
        if match_result['produto']:
            itens_identificados += 1

    await db.notas.update_one({'_id': ObjectId(nota_id)}, {'$set': {'itens_identificados': itens_identificados}})
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    items_docs = await db.itens_nota.find({'nota_id': nota_id}).to_list(1000)
    return {'nota': Nota.from_mongo(nota_doc).model_dump(), 'itens': [ItemNota.from_mongo(i).model_dump() for i in items_docs]}

@api_router.delete("/notas/{nota_id}")
async def delete_nota(nota_id: str):
    await db.notas.delete_one({'_id': ObjectId(nota_id)})
    await db.itens_nota.delete_many({'nota_id': nota_id})
    await db.historico_leituras.delete_many({'nota_id': nota_id})
    return {"ok": True}

# ── API: Conferencia ───────────────────────────────────────────────
@api_router.post("/conferencias/iniciar/{nota_id}")
async def iniciar_conferencia(nota_id: str):
    doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    if not doc:
        raise HTTPException(404, "Nota nao encontrada")

    # Block if any item has no binding
    pending = await db.itens_nota.count_documents({'nota_id': nota_id, 'produto_interno_id': None, 'ignorado': {'$ne': True}})
    if pending > 0:
        raise HTTPException(400, f"Existem {pending} produto(s) sem vinculo. Complete a vinculacao antes de iniciar a conferencia.")

    await db.notas.update_one(
        {'_id': ObjectId(nota_id)},
        {'$set': {'status': 'em_conferencia', 'conferencia_inicio': datetime.now(timezone.utc).isoformat()}}
    )
    await db.itens_nota.update_many({'nota_id': nota_id}, {'$set': {'quantidade_conferida': 0}})
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    items = await db.itens_nota.find({'nota_id': nota_id}).to_list(1000)
    return {'nota': Nota.from_mongo(nota_doc).model_dump(), 'itens': [ItemNota.from_mongo(i).model_dump() for i in items]}

# ── API: Quick Product Lookup ──────────────────────────────────────
@api_router.get("/produtos/buscar-codigo/{codigo}")
async def buscar_produto_por_codigo(codigo: str):
    """Quick lookup by internal code - for ENTER shortcut in binding screen."""
    prod = await db.produtos.find_one({'codigo': codigo, 'ativo': True})
    if not prod:
        raise HTTPException(404, "Produto nao encontrado com esse codigo")
    return Produto.from_mongo(prod).model_dump()

# ── API: Vinculacao (per-nota binding) ─────────────────────────────
@api_router.get("/vinculacao/{nota_id}")
async def get_vinculacao(nota_id: str):
    """Get nota info and items needing binding."""
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    if not nota_doc:
        raise HTTPException(404, "Nota nao encontrada")
    nota = Nota.from_mongo(nota_doc)
    all_items = await db.itens_nota.find({'nota_id': nota_id}).sort('numero_item', 1).to_list(1000)
    all_items_dumped = [ItemNota.from_mongo(i).model_dump() for i in all_items]
    pending = [i for i in all_items_dumped if not i.get('produto_interno_id') and not i.get('ignorado')]
    linked = [i for i in all_items_dumped if i.get('produto_interno_id')]
    return {
        'nota': nota.model_dump(),
        'itens': all_items_dumped,
        'pendentes': pending,
        'vinculados': linked,
        'total_pendentes': len(pending),
        'total_vinculados': len(linked),
        'todos_vinculados': len(pending) == 0,
    }

class LeituraInput(BaseModel):
    nota_id: str
    codigo_barras: str
    item_ativo_id: Optional[str] = None

def _item_game_payload(doc):
    o = ItemNota.from_mongo(doc)
    return {
        'id': o.id, 'cprod': o.cprod, 'ean': o.ean, 'descricao_nfe': o.descricao_nfe,
        'produto_interno_codigo': o.produto_interno_codigo,
        'produto_interno_descricao': o.produto_interno_descricao,
        'quantidade': o.quantidade, 'quantidade_conferida': o.quantidade_conferida,
        'unidade': o.unidade,
    }

async def _nota_progress(nota_id: str):
    all_items = await db.itens_nota.find({'nota_id': nota_id}).to_list(1000)
    completos = sum(1 for i in all_items if i.get('quantidade', 0) > 0 and i.get('quantidade_conferida', 0) >= i.get('quantidade', 0))
    return {'itens_completos': completos, 'total_itens': len(all_items)}

@api_router.post("/conferencias/leitura")
async def processar_leitura(data: LeituraInput):
    """Game-mode scan: each beep = 1 unit. Free order (scan opens the item).
    If an item is active, only its EAN counts; other codes return 'produto_errado'."""
    nota = await db.notas.find_one({'_id': ObjectId(data.nota_id)})
    if not nota:
        raise HTTPException(404, "Nota nao encontrada")

    codigo = data.codigo_barras.strip()

    # Resolve candidate items in this nota: by nota EAN, then by internal product EAN/code
    candidates = await db.itens_nota.find({'nota_id': data.nota_id, 'ean': codigo}).to_list(100)
    produto = await db.produtos.find_one({'ean': codigo, 'ativo': True})
    if not produto:
        produto = await db.produtos.find_one({'codigo': codigo, 'ativo': True})
    if produto:
        p = Produto.from_mongo(produto)
        seen = {str(c['_id']) for c in candidates}
        more = await db.itens_nota.find({'nota_id': data.nota_id, 'produto_interno_id': p.id}).to_list(100)
        candidates += [m for m in more if str(m['_id']) not in seen]

    progress = await _nota_progress(data.nota_id)

    if not candidates:
        if produto:
            p = Produto.from_mongo(produto)
            hist = HistoricoLeitura(nota_id=data.nota_id, codigo_barras=codigo, produto_interno_id=p.id, resultado='nao_pertence')
            await db.historico_leituras.insert_one(hist.to_mongo())
            return {'success': False, 'tipo': 'nao_pertence', 'message': 'Produto nao pertence a esta NF-e.',
                    'scanned': {'codigo': p.codigo, 'descricao': p.descricao}, 'item': None, 'progress': progress}
        hist = HistoricoLeitura(nota_id=data.nota_id, codigo_barras=codigo, resultado='desconhecido')
        await db.historico_leituras.insert_one(hist.to_mongo())
        return {'success': False, 'tipo': 'desconhecido', 'message': 'Codigo nao encontrado nesta nota.',
                'scanned': {'codigo': codigo}, 'item': None, 'progress': progress}

    # Determine target item
    item_doc = None
    if data.item_ativo_id:
        for c in candidates:
            if str(c['_id']) == data.item_ativo_id:
                item_doc = c
                break
        if item_doc is None:
            # Scanned a different product while another item is active and incomplete
            try:
                ativo = await db.itens_nota.find_one({'_id': ObjectId(data.item_ativo_id)})
            except Exception:
                ativo = None
            if ativo and ativo.get('quantidade_conferida', 0) < ativo.get('quantidade', 0):
                scanned_doc = candidates[0]
                hist = HistoricoLeitura(nota_id=data.nota_id, codigo_barras=codigo, item_nota_id=str(scanned_doc['_id']), resultado='produto_errado')
                await db.historico_leituras.insert_one(hist.to_mongo())
                return {'success': False, 'tipo': 'produto_errado', 'message': 'Produto nao e o indicado no visor!',
                        'scanned': _item_game_payload(scanned_doc), 'item': _item_game_payload(ativo), 'progress': progress}

    if item_doc is None:
        incomplete = [c for c in candidates if c.get('quantidade_conferida', 0) < c.get('quantidade', 0)]
        item_doc = incomplete[0] if incomplete else candidates[0]

    item_obj = ItemNota.from_mongo(item_doc)

    # Already fully counted
    if item_obj.quantidade_conferida >= item_obj.quantidade:
        hist = HistoricoLeitura(nota_id=data.nota_id, codigo_barras=codigo, item_nota_id=item_obj.id, resultado='ja_conferido')
        await db.historico_leituras.insert_one(hist.to_mongo())
        return {'success': False, 'tipo': 'ja_conferido', 'message': 'Este item ja foi conferido!',
                'item': _item_game_payload(item_doc), 'progress': progress}

    # Count +1
    new_qty = item_obj.quantidade_conferida + 1
    completed = new_qty >= item_obj.quantidade
    status = 'conferido' if completed else 'em_contagem'
    await db.itens_nota.update_one({'_id': ObjectId(item_obj.id)}, {'$set': {'quantidade_conferida': new_qty, 'status': status}})

    hist = HistoricoLeitura(nota_id=data.nota_id, codigo_barras=codigo, produto_interno_id=item_obj.produto_interno_id, item_nota_id=item_obj.id, resultado='encontrado')
    await db.historico_leituras.insert_one(hist.to_mongo())

    progress = await _nota_progress(data.nota_id)
    await db.notas.update_one({'_id': ObjectId(data.nota_id)}, {'$set': {'itens_conferidos': progress['itens_completos']}})

    payload = _item_game_payload(item_doc)
    payload['quantidade_conferida'] = new_qty
    return {
        'success': True,
        'tipo': 'completo' if completed else 'parcial',
        'message': 'Item conferido!' if completed else f'{new_qty:g} de {item_obj.quantidade:g}',
        'item': payload,
        'progress': progress,
        'nota_completa': progress['itens_completos'] >= progress['total_itens'],
    }

class ConfirmarVinculoInput(BaseModel):
    item_nota_id: str
    produto_interno_id: Optional[str] = None
    produto_interno_codigo: Optional[str] = None  # Allow binding by code
    origem_vinculo: str = "manual"

@api_router.post("/conferencias/confirmar-vinculo")
async def confirmar_vinculo(data: ConfirmarVinculoInput):
    item_doc = await db.itens_nota.find_one({'_id': ObjectId(data.item_nota_id)})
    if not item_doc:
        raise HTTPException(404, "Item nao encontrado")

    # Resolve product by ID or by internal code
    produto_doc = None
    if data.produto_interno_id:
        produto_doc = await db.produtos.find_one({'_id': ObjectId(data.produto_interno_id)})
    elif data.produto_interno_codigo:
        produto_doc = await db.produtos.find_one({'codigo': data.produto_interno_codigo, 'ativo': True})
    if not produto_doc:
        raise HTTPException(404, "Produto nao encontrado")

    item = ItemNota.from_mongo(item_doc)
    produto = Produto.from_mongo(produto_doc)
    nota_doc = await db.notas.find_one({'_id': ObjectId(item.nota_id)})
    nota = Nota.from_mongo(nota_doc) if nota_doc else None

    origem = data.origem_vinculo or 'manual'

    await db.itens_nota.update_one(
        {'_id': ObjectId(data.item_nota_id)},
        {'$set': {
            'produto_interno_id': produto.id,
            'produto_interno_codigo': produto.codigo,
            'produto_interno_descricao': produto.descricao,
            'metodo_identificacao': origem,
            'confianca': 100,
            'status': 'identificado',
        }}
    )

    # Save learning
    await save_learning(item, produto, nota, metodo=origem, confianca=100)

    items = await db.itens_nota.find({'nota_id': item.nota_id}).to_list(1000)
    identified = sum(1 for i in items if i.get('produto_interno_id'))
    await db.notas.update_one({'_id': ObjectId(item.nota_id)}, {'$set': {'itens_identificados': identified}})

    updated = await db.itens_nota.find_one({'_id': ObjectId(data.item_nota_id)})
    return ItemNota.from_mongo(updated).model_dump()

class JustificativaInput(BaseModel):
    item_nota_id: str
    justificativa: str

class FinalizarInput(BaseModel):
    operador: Optional[str] = None

class ConfirmarLoteInput(BaseModel):
    """Bulk binding confirmation."""
    vinculos: List[ConfirmarVinculoInput]

@api_router.post("/conferencias/justificativa")
async def salvar_justificativa(data: JustificativaInput):
    result = await db.itens_nota.update_one(
        {'_id': ObjectId(data.item_nota_id)},
        {'$set': {'justificativa_divergencia': data.justificativa}}
    )
    if result.matched_count == 0:
        raise HTTPException(404, "Item nao encontrado")
    return {"ok": True}

@api_router.post("/conferencias/confirmar-lote")
async def confirmar_vinculo_lote(data: ConfirmarLoteInput):
    """Bulk confirm multiple bindings at once."""
    confirmados = 0
    erros = []
    for v in data.vinculos:
        try:
            await confirmar_vinculo(v)
            confirmados += 1
        except Exception as e:
            erros.append({'item_nota_id': v.item_nota_id, 'erro': str(e)})
    return {'confirmados': confirmados, 'erros': erros, 'total': len(data.vinculos)}

@api_router.post("/conferencias/finalizar/{nota_id}")
async def finalizar_conferencia(nota_id: str, data: Optional[FinalizarInput] = None):
    items = await db.itens_nota.find({'nota_id': nota_id}).to_list(1000)
    has_divergence = any(i.get('quantidade_conferida', 0) > i.get('quantidade', 0) for i in items)
    has_pending = any(i.get('quantidade_conferida', 0) < i.get('quantidade', 0) for i in items)
    status = 'divergente' if (has_divergence or has_pending) else 'conferida'
    update = {'status': status, 'conferencia_fim': datetime.now(timezone.utc).isoformat()}
    if data and data.operador:
        update['operador_conferencia'] = data.operador.strip()
    await db.notas.update_one(
        {'_id': ObjectId(nota_id)},
        {'$set': update}
    )
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    return Nota.from_mongo(nota_doc).model_dump()

@api_router.get("/conferencias/relatorio/{nota_id}")
async def relatorio_conferencia(nota_id: str):
    """Detailed conference report data for A4 printing."""
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    if not nota_doc:
        raise HTTPException(404, "Nota nao encontrada")
    nota = Nota.from_mongo(nota_doc).model_dump()

    items_docs = await db.itens_nota.find({'nota_id': nota_id}).sort('numero_item', 1).to_list(1000)
    itens = []
    ok_count = falta_count = excedente_count = 0
    for d in items_docs:
        i = ItemNota.from_mongo(d).model_dump()
        qtd = i.get('quantidade', 0)
        conf = i.get('quantidade_conferida', 0)
        if conf == qtd:
            situacao = 'ok'
            ok_count += 1
        elif conf < qtd:
            situacao = 'falta'
            falta_count += 1
        else:
            situacao = 'excedente'
            excedente_count += 1
        itens.append({
            'numero_item': i.get('numero_item'),
            'cprod': i.get('cprod'),
            'ean': i.get('ean'),
            'descricao_nfe': i.get('descricao_nfe'),
            'produto_interno_codigo': i.get('produto_interno_codigo'),
            'produto_interno_descricao': i.get('produto_interno_descricao'),
            'unidade': i.get('unidade'),
            'quantidade': qtd,
            'quantidade_conferida': conf,
            'diferenca': conf - qtd,
            'situacao': situacao,
            'justificativa': i.get('justificativa_divergencia'),
        })

    # Duration
    duracao_min = None
    if nota.get('conferencia_inicio') and nota.get('conferencia_fim'):
        try:
            s = datetime.fromisoformat(nota['conferencia_inicio'])
            e = datetime.fromisoformat(nota['conferencia_fim'])
            duracao_min = round((e - s).total_seconds() / 60, 1)
        except Exception:
            pass

    total_leituras = await db.historico_leituras.count_documents({'nota_id': nota_id, 'resultado': 'encontrado'})

    return {
        'nota': nota,
        'itens': itens,
        'resumo': {
            'total_itens': len(itens),
            'itens_ok': ok_count,
            'itens_falta': falta_count,
            'itens_excedente': excedente_count,
            'total_leituras': total_leituras,
            'duracao_min': duracao_min,
            'resultado': 'CONFERIDA SEM DIVERGENCIAS' if (falta_count == 0 and excedente_count == 0) else 'CONFERIDA COM DIVERGENCIAS',
            'tudo_ok': falta_count == 0 and excedente_count == 0,
        },
    }

@api_router.post("/conferencias/relatorio/{nota_id}/salvar")
async def salvar_relatorio(nota_id: str):
    """Salva o snapshot do relatorio de conferencia dentro do documento da nota."""
    nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
    if not nota_doc:
        raise HTTPException(404, "Nota nao encontrada")

    # Reutiliza a mesma logica do relatorio para gerar o snapshot
    nota = Nota.from_mongo(nota_doc).model_dump()
    items_docs = await db.itens_nota.find({'nota_id': nota_id}).sort('numero_item', 1).to_list(1000)
    itens = []
    ok_count = falta_count = excedente_count = 0
    for d in items_docs:
        i = ItemNota.from_mongo(d).model_dump()
        qtd = i.get('quantidade', 0)
        conf = i.get('quantidade_conferida', 0)
        if conf == qtd:
            situacao = 'ok'; ok_count += 1
        elif conf < qtd:
            situacao = 'falta'; falta_count += 1
        else:
            situacao = 'excedente'; excedente_count += 1
        itens.append({
            'numero_item': i.get('numero_item'), 'cprod': i.get('cprod'),
            'descricao_nfe': i.get('descricao_nfe'),
            'produto_interno_codigo': i.get('produto_interno_codigo'),
            'produto_interno_descricao': i.get('produto_interno_descricao'),
            'unidade': i.get('unidade'), 'quantidade': qtd,
            'quantidade_conferida': conf, 'diferenca': conf - qtd,
            'situacao': situacao, 'justificativa': i.get('justificativa_divergencia'),
        })

    duracao_min = None
    if nota.get('conferencia_inicio') and nota.get('conferencia_fim'):
        try:
            s = datetime.fromisoformat(nota['conferencia_inicio'])
            e = datetime.fromisoformat(nota['conferencia_fim'])
            duracao_min = round((e - s).total_seconds() / 60, 1)
        except Exception:
            pass

    total_leituras = await db.historico_leituras.count_documents({'nota_id': nota_id, 'resultado': 'encontrado'})

    snapshot = {
        'nota': nota,
        'itens': itens,
        'resumo': {
            'total_itens': len(itens), 'itens_ok': ok_count,
            'itens_falta': falta_count, 'itens_excedente': excedente_count,
            'total_leituras': total_leituras, 'duracao_min': duracao_min,
            'resultado': 'CONFERIDA SEM DIVERGENCIAS' if (falta_count == 0 and excedente_count == 0) else 'CONFERIDA COM DIVERGENCIAS',
            'tudo_ok': falta_count == 0 and excedente_count == 0,
        },
        'salvo_em': datetime.now(timezone.utc).isoformat(),
    }

    await db.notas.update_one({'_id': ObjectId(nota_id)}, {'$set': {'relatorio_salvo': snapshot}})
    return {'ok': True, 'salvo_em': snapshot['salvo_em']}

# ── API: Recognition Center ────────────────────────────────────────
@api_router.get("/reconhecimento")
async def list_reconhecimento(fornecedor_cnpj: Optional[str] = None):
    """List all items without binding (pending recognition)."""
    query = {'produto_interno_id': None, 'ignorado': {'$ne': True}}
    if fornecedor_cnpj:
        nota_ids = []
        notas = await db.notas.find({'fornecedor_cnpj': fornecedor_cnpj}, {'_id': 1}).to_list(1000)
        nota_ids = [str(n['_id']) for n in notas]
        if nota_ids:
            query['nota_id'] = {'$in': nota_ids}
        else:
            return []

    items = await db.itens_nota.find(query).sort('nota_id', -1).to_list(1000)
    result = []
    nota_cache = {}
    for item_doc in items:
        item = ItemNota.from_mongo(item_doc)
        nota_id = item.nota_id
        if nota_id not in nota_cache:
            nota_doc = await db.notas.find_one({'_id': ObjectId(nota_id)})
            nota_cache[nota_id] = Nota.from_mongo(nota_doc) if nota_doc else None
        nota = nota_cache[nota_id]
        result.append({
            **item.model_dump(),
            'fornecedor_cnpj': nota.fornecedor_cnpj if nota else None,
            'fornecedor_nome': nota.fornecedor_nome if nota else None,
            'nota_numero': nota.numero if nota else None,
        })
    return result

@api_router.post("/reconhecimento/confirmar")
async def confirmar_reconhecimento(data: ConfirmarVinculoInput):
    """Confirm binding from Recognition Center - same as conference confirmar-vinculo."""
    return await confirmar_vinculo(data)

@api_router.post("/reconhecimento/ignorar/{item_id}")
async def ignorar_reconhecimento(item_id: str):
    """Mark item as ignored in recognition center."""
    result = await db.itens_nota.update_one(
        {'_id': ObjectId(item_id)},
        {'$set': {'ignorado': True}}
    )
    if result.matched_count == 0:
        raise HTTPException(404, "Item nao encontrado")
    return {"ok": True}

# ����� API: Equivalencias ─────────────────────────────────────────────
@api_router.get("/equivalencias")
async def list_equivalencias(fornecedor_cnpj: Optional[str] = None):
    query = {}
    if fornecedor_cnpj:
        query['fornecedor_cnpj'] = fornecedor_cnpj
    docs = await db.equivalencia_produtos.find(query).sort('created_at', -1).to_list(1000)
    return [EquivalenciaProduto.from_mongo(d).model_dump() for d in docs]

@api_router.delete("/equivalencias/{eq_id}")
async def delete_equivalencia(eq_id: str):
    await db.equivalencia_produtos.delete_one({'_id': ObjectId(eq_id)})
    return {"ok": True}

# ── API: Historico de Aprendizado ──────────────────────────────────
@api_router.get("/historico-aprendizado")
async def list_historico_aprendizado(limit: int = 50):
    docs = await db.historico_aprendizado.find().sort('created_at', -1).to_list(limit)
    return [HistoricoAprendizado.from_mongo(d).model_dump() for d in docs]

# ── API: Dashboard Inteligente ─────────────────────────────────────
@api_router.get("/dashboard")
async def dashboard():
    today_start = datetime.combine(date.today(), datetime.min.time()).replace(tzinfo=timezone.utc).isoformat()
    total_notas = await db.notas.count_documents({})
    conferidas_hoje = await db.notas.count_documents({'status': 'conferida', 'conferencia_fim': {'$gte': today_start}})
    pendentes = await db.notas.count_documents({'status': {'$in': ['pendente', 'em_conferencia']}})
    divergentes = await db.notas.count_documents({'status': 'divergente'})

    conferidas = await db.notas.find({'status': 'conferida', 'conferencia_inicio': {'$ne': None}, 'conferencia_fim': {'$ne': None}}).to_list(100)
    avg_time = 0
    if conferidas:
        times = []
        for n in conferidas:
            try:
                s = datetime.fromisoformat(n['conferencia_inicio'])
                e = datetime.fromisoformat(n['conferencia_fim'])
                times.append((e - s).total_seconds())
            except Exception:
                pass
        if times:
            avg_time = round(sum(times) / len(times) / 60, 1)

    total_itens = await db.itens_nota.count_documents({})
    sem_vinculo = await db.itens_nota.count_documents({'produto_interno_id': None, 'ignorado': {'$ne': True}})
    total_matched = await db.itens_nota.count_documents({'produto_interno_id': {'$ne': None}})

    # Recognition stats by method
    by_ean = await db.itens_nota.count_documents({'metodo_identificacao': 'ean'})
    by_vinculo = await db.itens_nota.count_documents({'metodo_identificacao': 'vinculo'})
    by_similaridade = await db.itens_nota.count_documents({'metodo_identificacao': {'$in': ['similaridade', 'vinculo_aprendido']}})
    by_manual = await db.itens_nota.count_documents({'metodo_identificacao': 'manual'})
    auto_matched = by_ean + by_vinculo + by_similaridade

    precisao = round((auto_matched / total_matched * 100) if total_matched > 0 else 0, 1)
    pct_auto = round((auto_matched / total_itens * 100) if total_itens > 0 else 0, 1)
    total_equivalencias = await db.equivalencia_produtos.count_documents({})
    total_aprendizado = await db.historico_aprendizado.count_documents({})

    pipeline = [{'$group': {'_id': '$fornecedor_nome', 'count': {'$sum': 1}}}, {'$sort': {'count': -1}}, {'$limit': 5}]
    top_fornecedores = await db.notas.aggregate(pipeline).to_list(5)

    # Suppliers with highest error rates (most unmatched items) — single aggregation
    # replaces the previous O(N) loop that fired 2 count_documents per nota.
    fornecedor_errors = []
    error_pipeline = [
        {'$lookup': {
            'from': 'itens_nota',
            'let': {'nid': {'$toString': '$_id'}},
            'pipeline': [
                {'$match': {'$expr': {'$eq': ['$nota_id', '$$nid']}}},
                {'$group': {
                    '_id': None,
                    'total': {'$sum': 1},
                    'unmatched': {'$sum': {'$cond': [{'$eq': ['$produto_interno_id', None]}, 1, 0]}},
                }},
            ],
            'as': 'stats',
        }},
        {'$unwind': {'path': '$stats', 'preserveNullAndEmptyArrays': True}},
        {'$group': {
            '_id': {'$ifNull': ['$fornecedor_nome', 'N/A']},
            'total': {'$sum': {'$ifNull': ['$stats.total', 0]}},
            'unmatched': {'$sum': {'$ifNull': ['$stats.unmatched', 0]}},
        }},
    ]
    async for f in db.notas.aggregate(error_pipeline):
        total = f.get('total', 0)
        if total > 0:
            unmatched = f.get('unmatched', 0)
            fornecedor_errors.append({
                'nome': f.get('_id') or 'N/A',
                'taxa_erro': round(unmatched / total * 100, 1),
                'total': total,
                'sem_vinculo': unmatched,
            })
    fornecedor_errors.sort(key=lambda x: x['taxa_erro'], reverse=True)

    notas_por_dia = []
    for i in range(6, -1, -1):
        d = date.today() - timedelta(days=i)
        ds = datetime.combine(d, datetime.min.time()).replace(tzinfo=timezone.utc).isoformat()
        de = datetime.combine(d, datetime.max.time()).replace(tzinfo=timezone.utc).isoformat()
        count = await db.notas.count_documents({'created_at': {'$gte': ds, '$lte': de}})
        notas_por_dia.append({'data': d.isoformat(), 'quantidade': count})

    return {
        'total_notas': total_notas, 'conferidas_hoje': conferidas_hoje, 'pendentes': pendentes,
        'divergentes': divergentes, 'tempo_medio_min': avg_time, 'sem_vinculo': sem_vinculo,
        'precisao_reconhecimento': precisao, 'pct_identificacao_auto': pct_auto,
        'total_equivalencias': total_equivalencias, 'total_aprendizado': total_aprendizado,
        'reconhecimento_por_metodo': {
            'ean': by_ean, 'vinculo': by_vinculo,
            'similaridade': by_similaridade, 'manual': by_manual,
        },
        'fornecedor_errors': fornecedor_errors[:5],
        'top_fornecedores': [{'nome': f['_id'] or 'N/A', 'quantidade': f['count']} for f in top_fornecedores],
        'notas_por_dia': notas_por_dia,
        'total_itens': total_itens, 'total_matched': total_matched,
    }

# ── API: Seed Data ─────────────────────────────────────────────────
@api_router.post("/seed")
async def seed_data():
    await db.produtos.delete_many({})
    await db.fornecedores.delete_many({})
    await db.equivalencia_produtos.delete_many({})
    await db.notas.delete_many({})
    await db.itens_nota.delete_many({})
    await db.historico_leituras.delete_many({})
    await db.historico_aprendizado.delete_many({})

    produtos = [
        {"codigo": "10001", "descricao": "AMORTECEDOR DIANTEIRO GOL G5", "ean": "7891234567001", "unidade": "UN", "preco": 189.90, "categoria": "Suspensao"},
        {"codigo": "10002", "descricao": "AMORTECEDOR DIANTEIRO GOL G6", "ean": "7891234567002", "unidade": "UN", "preco": 199.90, "categoria": "Suspensao"},
        {"codigo": "10003", "descricao": "AMORTECEDOR TRASEIRO GOL G5", "ean": "7891234567003", "unidade": "UN", "preco": 149.90, "categoria": "Suspensao"},
        {"codigo": "10004", "descricao": "AMORTECEDOR DIANTEIRO SAVEIRO G5", "ean": "7891234567004", "unidade": "UN", "preco": 209.90, "categoria": "Suspensao"},
        {"codigo": "10005", "descricao": "PASTILHA DE FREIO DIANTEIRA GOL G5", "ean": "7891234567005", "unidade": "JG", "preco": 89.90, "categoria": "Freios"},
        {"codigo": "10006", "descricao": "PASTILHA DE FREIO DIANTEIRA GOL G6", "ean": "7891234567006", "unidade": "JG", "preco": 94.90, "categoria": "Freios"},
        {"codigo": "10007", "descricao": "DISCO DE FREIO DIANTEIRO GOL G5", "ean": "7891234567007", "unidade": "UN", "preco": 119.90, "categoria": "Freios"},
        {"codigo": "10008", "descricao": "FILTRO DE OLEO MOTOR GOL 1.0", "ean": "7891234567008", "unidade": "UN", "preco": 24.90, "categoria": "Filtros"},
        {"codigo": "10009", "descricao": "FILTRO DE AR MOTOR GOL 1.0", "ean": "7891234567009", "unidade": "UN", "preco": 34.90, "categoria": "Filtros"},
        {"codigo": "10010", "descricao": "FILTRO DE COMBUSTIVEL GOL 1.0", "ean": "7891234567010", "unidade": "UN", "preco": 29.90, "categoria": "Filtros"},
        {"codigo": "10011", "descricao": "CORREIA DENTADA GOL 1.0", "ean": "7891234567011", "unidade": "UN", "preco": 49.90, "categoria": "Motor"},
        {"codigo": "10012", "descricao": "CORREIA ALTERNADOR GOL 1.0", "ean": "7891234567012", "unidade": "UN", "preco": 34.90, "categoria": "Motor"},
        {"codigo": "10013", "descricao": "VELA DE IGNICAO GOL 1.0", "ean": "7891234567013", "unidade": "UN", "preco": 19.90, "categoria": "Motor"},
        {"codigo": "10014", "descricao": "ROLAMENTO RODA DIANTEIRA GOL G5", "ean": "7891234567014", "unidade": "UN", "preco": 79.90, "categoria": "Suspensao"},
        {"codigo": "10015", "descricao": "TERMINAL DIRECAO GOL G5", "ean": "7891234567015", "unidade": "UN", "preco": 54.90, "categoria": "Direcao"},
        {"codigo": "10016", "descricao": "BIELETA BARRA ESTABILIZADORA GOL G5", "ean": "7891234567016", "unidade": "UN", "preco": 39.90, "categoria": "Suspensao"},
        {"codigo": "10017", "descricao": "PIVO SUSPENSAO GOL G5", "ean": "7891234567017", "unidade": "UN", "preco": 69.90, "categoria": "Suspensao"},
        {"codigo": "10018", "descricao": "BUCHA BANDEJA GOL G5", "ean": "7891234567018", "unidade": "UN", "preco": 29.90, "categoria": "Suspensao"},
        {"codigo": "10019", "descricao": "BOMBA DAGUA GOL 1.0", "ean": "7891234567019", "unidade": "UN", "preco": 89.90, "categoria": "Motor"},
        {"codigo": "10020", "descricao": "JUNTA CABECOTE GOL 1.0", "ean": "7891234567020", "unidade": "UN", "preco": 159.90, "categoria": "Motor"},
        # Heavy truck parts (matching NF-e items from BR COMPANY)
        {"codigo": "20001", "descricao": "BUJAO CARTER COM IMA 26MM", "ean": None, "unidade": "PC", "preco": 6.50, "categoria": "Motor"},
        {"codigo": "20002", "descricao": "CAPA COLUNA DIRECAO MERCEDES", "ean": None, "unidade": "PC", "preco": 42.00, "categoria": "Direcao"},
        {"codigo": "20003", "descricao": "TERMINAL ACELERADOR COMPLETO 6X1", "ean": None, "unidade": "PC", "preco": 9.50, "categoria": "Motor"},
        {"codigo": "20004", "descricao": "MOLA FECHADURA PORTA", "ean": None, "unidade": "PC", "preco": 0.85, "categoria": "Carroceria"},
        {"codigo": "20005", "descricao": "CONEXAO REDUTORA MACHO FEMEA M22X1.5", "ean": None, "unidade": "PC", "preco": 4.80, "categoria": "Conexoes"},
        {"codigo": "20006", "descricao": "RESERVATORIO AR 40 LITROS CHAPA", "ean": None, "unidade": "PC", "preco": 280.00, "categoria": "Freios"},
        {"codigo": "20007", "descricao": "ESTICADOR CORREIA POLIA ACO", "ean": None, "unidade": "PC", "preco": 155.00, "categoria": "Motor"},
        {"codigo": "20008", "descricao": "CUBO EMBREAGEM COM ROLAMENTO 366", "ean": None, "unidade": "PC", "preco": 120.00, "categoria": "Embreagem"},
        {"codigo": "20009", "descricao": "FILTRO AR COMPLETO PLASTICO WORKER", "ean": None, "unidade": "PC", "preco": 265.00, "categoria": "Filtros"},
        {"codigo": "20010", "descricao": "PINO CILINDRICO RODA DIANTEIRO TRASEIRO", "ean": None, "unidade": "PC", "preco": 7.20, "categoria": "Rodas"},
        {"codigo": "20011", "descricao": "JOGO REPARO TRANSFERENCIA CAMBIO", "ean": None, "unidade": "JG", "preco": 38.00, "categoria": "Cambio"},
        {"codigo": "20012", "descricao": "TAMPA TANQUE COMBUSTIVEL COM ROSCA ACO", "ean": None, "unidade": "PC", "preco": 35.00, "categoria": "Combustivel"},
        {"codigo": "20013", "descricao": "TOMADA DE AR CARROCERIA", "ean": None, "unidade": "PC", "preco": 115.00, "categoria": "Carroceria"},
        {"codigo": "20014", "descricao": "ESTICADOR MOLA ESPIRAL MOTOR CUMMINS", "ean": None, "unidade": "PC", "preco": 125.00, "categoria": "Motor"},
        {"codigo": "20015", "descricao": "SAPATA FREIO EIXO TRASEIRO 13X3", "ean": None, "unidade": "PC", "preco": 65.00, "categoria": "Freios"},
        {"codigo": "20016", "descricao": "MANCAL EMBREAGEM COM ROLAMENTO 1 LINHA", "ean": None, "unidade": "PC", "preco": 195.00, "categoria": "Embreagem"},
        {"codigo": "20017", "descricao": "CONEXAO BOMBA ARLA TUBO 8MM", "ean": None, "unidade": "PC", "preco": 10.00, "categoria": "Conexoes"},
        {"codigo": "20018", "descricao": "JUNTA TAMPA DISTRIBUICAO CUMMINS 6 CILINDROS", "ean": None, "unidade": "PC", "preco": 30.00, "categoria": "Motor"},
        {"codigo": "20019", "descricao": "MANGUEIRA ENCHIMENTO OLEO MOTOR", "ean": None, "unidade": "PC", "preco": 90.00, "categoria": "Motor"},
        {"codigo": "20020", "descricao": "CONEXAO UNIAO REDUTORA PLASTICA 8X6MM", "ean": None, "unidade": "PC", "preco": 1.10, "categoria": "Conexoes"},
        {"codigo": "20021", "descricao": "MANGUEIRA INFERIOR RADIADOR", "ean": None, "unidade": "PC", "preco": 160.00, "categoria": "Arrefecimento"},
    ]
    for p in produtos:
        prod = Produto(**p)
        await db.produtos.insert_one(prod.to_mongo())

    fornecedores = [
        {"cnpj": "61099008000141", "nome": "COFAP AUTOPECAS LTDA", "contato": "vendas@cofap.com.br", "telefone": "(11) 4174-8100"},
        {"cnpj": "88611835000129", "nome": "FRAS-LE S.A.", "contato": "vendas@frasle.com", "telefone": "(54) 3289-1500"},
        {"cnpj": "59104273000191", "nome": "MANN-FILTER BRASIL", "contato": "vendas@mann-filter.com.br", "telefone": "(11) 4166-7000"},
        {"cnpj": "02982122000130", "nome": "GATES DO BRASIL LTDA", "contato": "vendas@gates.com.br", "telefone": "(11) 3611-1000"},
        {"cnpj": "10313421000126", "nome": "BR COMPANY IMP E EXP DE PECAS LTDA", "contato": "vendas@brcompany.com.br", "telefone": "(11) 4092-3050"},
    ]
    for f in fornecedores:
        forn = Fornecedor(**f)
        await db.fornecedores.insert_one(forn.to_mongo())

    for eq_data in [
        {"fornecedor_cnpj": "61099008000141", "fornecedor_nome": "COFAP AUTOPECAS LTDA", "codigo_fornecedor": "GP32918", "ean": "7891234567001", "descricao_nfe": "AMORT. DT GOL G5", "produto_interno_codigo": "10001", "produto_interno_descricao": "AMORTECEDOR DIANTEIRO GOL G5"},
        {"fornecedor_cnpj": "88611835000129", "fornecedor_nome": "FRAS-LE S.A.", "codigo_fornecedor": "PD/1040", "ean": "7891234567005", "descricao_nfe": "PAST. FREIO DT GOL G5", "produto_interno_codigo": "10005", "produto_interno_descricao": "PASTILHA DE FREIO DIANTEIRA GOL G5"},
    ]:
        prod = await db.produtos.find_one({'codigo': eq_data['produto_interno_codigo']})
        if prod:
            eq_data['produto_interno_id'] = str(prod['_id'])
            eq = EquivalenciaProduto(**eq_data)
            await db.equivalencia_produtos.insert_one(eq.to_mongo())

    return {"message": "Dados de exemplo criados com sucesso!", "produtos": len(produtos), "fornecedores": len(fornecedores)}

@api_router.get("/sample-xml")
async def get_sample_xml():
    xml = '''<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00">
  <NFe xmlns="http://www.portalfiscal.inf.br/nfe">
    <infNFe versao="4.00" Id="NFe35240112345678000190550010000012341234567890">
      <ide><cUF>35</cUF><cNF>12345678</cNF><natOp>VENDA DE MERCADORIAS</natOp><mod>55</mod><serie>1</serie><nNF>1234</nNF><dhEmi>2024-01-15T10:30:00-03:00</dhEmi><tpNF>1</tpNF></ide>
      <emit><CNPJ>61099008000141</CNPJ><xNome>COFAP AUTOPECAS LTDA</xNome><enderEmit><xLgr>RUA EXEMPLO</xLgr><nro>100</nro><xBairro>CENTRO</xBairro><cMun>3550308</cMun><xMun>SAO PAULO</xMun><UF>SP</UF><CEP>01001000</CEP></enderEmit><IE>123456789</IE></emit>
      <det nItem="1"><prod><cProd>GP32918</cProd><cEAN>7891234567001</cEAN><xProd>AMORT. DT GOL G5</xProd><NCM>87083090</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>4.0000</qCom><vUnCom>189.9000</vUnCom><vProd>759.60</vProd><cEANTrib>7891234567001</cEANTrib><uTrib>UN</uTrib><qTrib>4.0000</qTrib><vUnTrib>189.9000</vUnTrib></prod></det>
      <det nItem="2"><prod><cProd>GP32919</cProd><cEAN>7891234567003</cEAN><xProd>AMORT. TRAS. GOL G5</xProd><NCM>87083090</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>4.0000</qCom><vUnCom>149.9000</vUnCom><vProd>599.60</vProd><cEANTrib>7891234567003</cEANTrib><uTrib>UN</uTrib><qTrib>4.0000</qTrib><vUnTrib>149.9000</vUnTrib></prod></det>
      <det nItem="3"><prod><cProd>GP45001</cProd><cEAN>SEM GTIN</cEAN><xProd>AMORT. DT SAVEIRO G5</xProd><NCM>87083090</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>2.0000</qCom><vUnCom>209.9000</vUnCom><vProd>419.80</vProd><cEANTrib>SEM GTIN</cEANTrib><uTrib>UN</uTrib><qTrib>2.0000</qTrib><vUnTrib>209.9000</vUnTrib></prod></det>
      <det nItem="4"><prod><cProd>GP99001</cProd><cEAN>SEM GTIN</cEAN><xProd>BUCHA BANDEJA DT GOL G5</xProd><NCM>87083090</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>4.0000</qCom><vUnCom>29.9000</vUnCom><vProd>119.60</vProd><cEANTrib>SEM GTIN</cEANTrib><uTrib>UN</uTrib><qTrib>4.0000</qTrib><vUnTrib>29.9000</vUnTrib></prod></det>
      <total><ICMSTot><vBC>0.00</vBC><vICMS>0.00</vICMS><vProd>1898.60</vProd><vNF>1898.60</vNF></ICMSTot></total>
    </infNFe>
  </NFe>
</nfeProc>'''
    return {"xml": xml}

# ── App Config ─────────────────────────────────────────────────────
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def startup():
    await db.produtos.create_index("codigo")
    await db.produtos.create_index("ean")
    await db.produtos.create_index("ativo")
    await db.fornecedores.create_index("cnpj")
    await db.itens_nota.create_index("nota_id")
    await db.itens_nota.create_index([("produto_interno_id", 1)])
    await db.itens_nota.create_index([("cprod", 1), ("produto_interno_codigo", 1)])
    await db.itens_nota.create_index("metodo_identificacao")
    await db.itens_nota.create_index("ignorado")
    await db.notas.create_index("chave")
    await db.notas.create_index("status")
    await db.notas.create_index("fornecedor_cnpj")
    await db.notas.create_index([("created_at", -1)])
    await db.notas.create_index("conferencia_fim")
    await db.equivalencia_produtos.create_index([("fornecedor_cnpj", 1), ("codigo_fornecedor", 1)])
    await db.historico_leituras.create_index("nota_id")
    await db.historico_aprendizado.create_index("fornecedor_cnpj")
    await db.historico_aprendizado.create_index("created_at")
    logger.info("NF-e Conference System v2.0 started - Intelligent Matching Engine")

@app.on_event("shutdown")
async def shutdown():
    client.close()
